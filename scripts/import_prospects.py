"""
Clean & normalise the marketing FET prospects workbook into a tidy JSON the
backend imports. Handles the sheet quirks: inconsistent/"HOTEL" headers, the
Manufacturing title row, the empty sheet, column drift (phones landing in the
email column), malformed emails, and duplicates.

Run:  python scripts/import_prospects.py
Out:  backend/database/data/fet-prospects.json
Then: php artisan prospects:import   (in backend/)
"""

import os, re, json
import openpyxl

ROOT = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
SRC = os.path.join(ROOT, "FUEL ECO TEC PROSPECTS.xlsx")
OUT_DIR = os.path.join(ROOT, "backend", "database", "data")
OUT = os.path.join(OUT_DIR, "fet-prospects.json")
os.makedirs(OUT_DIR, exist_ok=True)

SOURCE = "marketing import 2026-06"

# Sheet title -> canonical category
CATEGORY = {
    "CARGO CO.": "CARGO",
    "DISTRIBUTERS": "DISTRIBUTOR",
    "CONSTRUCTION CO.": "CONSTRUCTION",
    "MANUFACTURING PLANTS": "MANUFACTURING",
    "PUBLIC TP CO.S": "PUBLIC_TRANSPORT",
    "SCHOOLS": "SCHOOL",
    "FARMERS": "FARMER",
    "SPARE PARTS SHOPS & GARAGES": "SPARE_PARTS",
    "CAR BONDS": "CAR_BOND",
    "FUNERAL SERVICES": "FUNERAL",
}

EMAIL_RE = re.compile(r"^[^@\s]+@[^@\s]+\.[^@\s]+$")


def digits(s: str) -> str:
    return re.sub(r"\D", "", s)


def is_phone(v: str) -> bool:
    return "@" not in v and len(digits(v)) >= 7


def clean_email(v: str):
    """Return (email_or_None, ok). Strips spaces; drops obvious placeholders."""
    e = v.strip().replace(" ", "").lower()
    if "example.com" in e or e in ("", "-"):
        return None, False
    return e, bool(EMAIL_RE.match(e))


def header_index(rows):
    """Row index whose cells include 'location' (handles title rows / HOTEL headers)."""
    for i, r in enumerate(rows[:4]):
        if any(c and "location" in str(c).lower() for c in r):
            return i
    return 0


def col_index(header, *needles):
    for i, c in enumerate(header):
        if c and any(n in str(c).lower() for n in needles):
            return i
    return None


def map_status(status_raw: str, feedback_raw: str) -> str:
    s = (status_raw or "").strip().lower()
    f = (feedback_raw or "").strip().lower()
    if "not delivered" in f or "bounce" in f or "fail" in f:
        return "bounced"
    if s in ("sent", "delivered", "contacted"):
        return "contacted"
    return "not_contacted"


def parse():
    wb = openpyxl.load_workbook(SRC, data_only=True)
    out = []
    seen = set()           # (category, name_lower) — dedupe / unique key
    stats = {"total": 0, "bad_email": 0, "no_contact": 0, "dupes": 0, "by_cat": {}}

    for ws in wb.worksheets:
        category = CATEGORY.get(ws.title.strip(), ws.title.strip().upper().replace(" ", "_"))
        rows = [r for r in ws.iter_rows(values_only=True)]
        if not rows:
            continue
        hi = header_index(rows)
        header = rows[hi] if hi < len(rows) else ()
        s_idx = col_index(header, "status")
        f_idx = col_index(header, "feedback")
        fu_idx = col_index(header, "followup", "follow up", "follow-up")

        for raw in rows[hi + 1:]:
            if not raw or not raw[0] or not str(raw[0]).strip():
                continue
            name = str(raw[0]).strip()
            key = (category, name.lower())
            if key in seen:
                stats["dupes"] += 1
                continue
            seen.add(key)

            cells = [str(c).strip() for c in raw[1:] if c is not None and str(c).strip()]
            emails = [c for c in cells if "@" in c]
            phones = [c for c in cells if is_phone(c)]
            texts = [c for c in cells if "@" not in c and not is_phone(c)]

            email, ok = (None, True)
            flags = []
            if emails:
                email, ok = clean_email(emails[0])
                if email and not ok:
                    flags.append("bad_email")
                if emails[0] and email is None:
                    flags.append("bad_email")

            phone = " / ".join(dict.fromkeys(phones)) or None      # de-dupe, keep order
            location = texts[0] if texts else None

            status_raw = raw[s_idx] if s_idx is not None and s_idx < len(raw) else None
            fb_raw = raw[f_idx] if f_idx is not None and f_idx < len(raw) else None
            fu_raw = raw[fu_idx] if fu_idx is not None and fu_idx < len(raw) else None
            status = map_status(str(status_raw or ""), str(fb_raw or ""))

            if not email and not phone:
                flags.append("no_contact")
                stats["no_contact"] += 1
            if "bad_email" in flags:
                stats["bad_email"] += 1

            out.append({
                "name": name,
                "category": category,
                "product": "FET",
                "location": location,
                "phone": phone,
                "email": email,
                "outreach_status": status,
                "feedback": (str(fb_raw).strip() if fb_raw else None),
                "follow_up": (str(fu_raw).strip() if fu_raw else None),
                "flags": flags or None,
                "source": SOURCE,
            })
            stats["total"] += 1
            stats["by_cat"][category] = stats["by_cat"].get(category, 0) + 1

    with open(OUT, "w", encoding="utf-8") as f:
        json.dump(out, f, ensure_ascii=False, indent=2)

    print(f"wrote {OUT}")
    print(f"  total={stats['total']}  bad_email={stats['bad_email']}  "
          f"no_contact={stats['no_contact']}  dupes_skipped={stats['dupes']}")
    for cat, n in sorted(stats["by_cat"].items()):
        print(f"  {cat:18} {n}")


if __name__ == "__main__":
    parse()
